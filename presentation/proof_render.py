"""Приблизительный рендер .pptx в PNG для визуальной проверки макета.

Не заменяет PowerPoint: шрифты DejaVu вместо Arial, без градиентов и теней.
Нужен только для контроля переполнений и наложений.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR

import matplotlib
FONTS = Path(matplotlib.get_data_path()) / 'fonts' / 'ttf'
REG = FONTS / 'DejaVuSans.ttf'
BOLD = FONTS / 'DejaVuSans-Bold.ttf'
MONO = FONTS / 'DejaVuSansMono.ttf'
DPI = 110
_cache: dict = {}


def font(px, bold=False, mono=False):
    key = (px, bold, mono)
    if key not in _cache:
        path = MONO if mono else (BOLD if bold else REG)
        _cache[key] = ImageFont.truetype(str(path), max(px, 6))
    return _cache[key]


def wrap(lines, fnt, width_px):
    out = []
    for line in lines:
        if not line:
            out.append('')
            continue
        words = line.split(' ')
        cur = ''
        for w in words:
            trial = (cur + ' ' + w).strip() if cur else w
            if fnt.getlength(trial) <= width_px or not cur:
                cur = trial
            else:
                out.append(cur)
                cur = w
        out.append(cur)
    return out


def draw_text(dr, shape, W, H):
    tf = shape.text_frame
    x = shape.left / 914400 * DPI
    y = shape.top / 914400 * DPI
    w = shape.width / 914400 * DPI
    h = shape.height / 914400 * DPI
    blocks = []
    for p in tf.paragraphs:
        runs = [(r.text, r) for r in p.runs]
        if not runs:
            continue
        size = max((r.font.size.pt if r.font.size else 14) for _, r in runs)
        color = next((r.font.color.rgb for _, r in runs
                      if r.font.color and r.font.color.type is not None), None)
        bold = any(r.font.bold for _, r in runs)
        mono = any((r.font.name or '').lower().startswith('consola') for _, r in runs)
        text_ = ''.join(t for t, _ in runs)
        ls = p.line_spacing or 1.15
        blocks.append((text_, size, color, bold, mono, ls, p.alignment))
    if not blocks:
        return
    total_h = 0
    rendered = []
    for text_, size, color, bold, mono, ls, al in blocks:
        px = size * DPI / 72
        fnt = font(int(px), bold, mono)
        lines = wrap(text_.split('\n'), fnt, max(w - 2, 10))
        lh = px * (ls if ls else 1.15) * 1.18
        rendered.append((lines, fnt, color, lh, al))
        total_h += lh * len(lines)
    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        y += (h - total_h) / 2
    elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
        y += h - total_h
    for lines, fnt, color, lh, al in rendered:
        for line in lines:
            tw = fnt.getlength(line)
            tx = x
            if al is not None and int(al) == 2:      # center
                tx = x + (w - tw) / 2
            elif al is not None and int(al) == 3:    # right
                tx = x + w - tw
            rgb = tuple(color) if color is not None else (255, 255, 255)
            dr.text((tx, y), line, font=fnt, fill=rgb)
            y += lh


def fill_rgb(shape):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            c = shape.fill.fore_color.rgb
            return tuple(c)
    except Exception:
        pass
    return None


def render(pptx_path, out_dir='presentation/proof'):
    prs = Presentation(pptx_path)
    W = int(prs.slide_width / 914400 * DPI)
    H = int(prs.slide_height / 914400 * DPI)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new('RGB', (W, H), (11, 18, 32))
        dr = ImageDraw.Draw(img)
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pic = Image.open(shape.image.blob and __import__('io').BytesIO(shape.image.blob))
                pic = pic.convert('RGB')
                tw = int(shape.width / 914400 * DPI)
                th = int(shape.height / 914400 * DPI)
                pic = pic.resize((max(tw, 2), max(th, 2)))
                img.paste(pic, (int(shape.left / 914400 * DPI), int(shape.top / 914400 * DPI)))
                continue
            if shape.has_table:
                tbl = shape.table
                x = shape.left / 914400 * DPI
                y = shape.top / 914400 * DPI
                for r in range(len(tbl.rows)):
                    rh = tbl.rows[r].height / 914400 * DPI
                    x = shape.left / 914400 * DPI
                    for c in range(len(tbl.columns)):
                        cw = tbl.columns[c].width / 914400 * DPI
                        cell = tbl.cell(r, c)
                        col = None
                        try:
                            col = tuple(cell.fill.fore_color.rgb)
                        except Exception:
                            col = (17, 28, 48)
                        dr.rectangle([x, y, x + cw, y + rh], fill=col)
                        txt = cell.text
                        if txt:
                            run = cell.text_frame.paragraphs[0].runs[0]
                            px = int((run.font.size.pt if run.font.size else 11) * DPI / 72)
                            fnt = font(px, bool(run.font.bold))
                            rgb = tuple(run.font.color.rgb) if run.font.color and \
                                run.font.color.type is not None else (234, 242, 255)
                            tw = fnt.getlength(txt)
                            tx = x + (cw - tw) / 2 if c > 0 else x + 0.10 * DPI
                            dr.text((tx, y + rh / 2 - px * 0.7), txt, font=fnt, fill=rgb)
                        x += cw
                    y += rh
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                dr_ = draw_text(dr, shape, W, H)
            elif fill_rgb(shape):
                x = shape.left / 914400 * DPI
                y = shape.top / 914400 * DPI
                w = shape.width / 914400 * DPI
                h = shape.height / 914400 * DPI
                dr.rectangle([x, y, x + w, y + h], fill=fill_rgb(shape))
            else:
                x = shape.left / 914400 * DPI
                y = shape.top / 914400 * DPI
                w = shape.width / 914400 * DPI
                h = shape.height / 914400 * DPI
                try:
                    if shape.line.color and shape.line.color.type is not None:
                        dr.rectangle([x, y, x + w, y + h],
                                     outline=tuple(shape.line.color.rgb))
                except Exception:
                    pass
        p = out / f'slide_{i:02d}.png'
        img.save(p)
        paths.append(p)
    return paths


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else 'presentation/PINN_EOS_Al.pptx'
    ps = render(src)
    print(f'rendered {len(ps)} slides -> {ps[0].parent}')
