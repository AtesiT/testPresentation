"""Программная проверка макета: переполнение текстовых блоков, выход за холст, наложения."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import proof_render as PR
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

W, H = 13.333, 7.5


def boxes(prs):
    for i, sl in enumerate(prs.slides, 1):
        out = []
        for sh in sl.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or sh.has_table:
                continue
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            x, y = sh.left / 914400, sh.top / 914400
            w, h = sh.width / 914400, sh.height / 914400
            need = 0.0
            for p in sh.text_frame.paragraphs:
                runs = [r for r in p.runs if r.text]
                if not runs:
                    continue
                size = max((r.font.size.pt if r.font.size else 14) for r in runs)
                text_ = ''.join(r.text for r in runs)
                mono = any((r.font.name or '').lower().startswith('consola') for r in runs)
                fnt = PR.font(int(size * PR.DPI / 72),
                              any(r.font.bold for r in runs), mono)
                lines = PR.wrap(text_.split('\n'), fnt, max(w * PR.DPI - 4, 20))
                ls = p.line_spacing or 1.15
                need += len(lines) * size * ls * 1.18 / 72
            out.append(dict(idx=i, text=sh.text_frame.text[:60].replace('\n', ' / '),
                            x=x, y=y, w=w, h=h, need=need))
        yield out


def report(path, tol=1.06, slack=0.03):
    prs = Presentation(path)
    issues = []
    for bs in boxes(prs):
        for b in bs:
            if b['need'] > b['h'] * tol + slack:
                issues.append(f"стр.{b['idx']:>2} переполнение {b['need']:.2f} > {b['h']:.2f} in · {b['text']}")
            if b['x'] < -0.02 or b['y'] < -0.02 or b['x'] + b['w'] > W + 0.02 or b['y'] + b['h'] > H + 0.02:
                issues.append(f"стр.{b['idx']:>2} выход за холст ({b['x']:.2f},{b['y']:.2f},"
                              f"{b['w']:.2f}×{b['h']:.2f}) · {b['text']}")
        for i in range(len(bs)):
            for j in range(i + 1, len(bs)):
                a, b = bs[i], bs[j]
                ox = min(a['x'] + a['w'], b['x'] + b['w']) - max(a['x'], b['x'])
                oy = min(a['y'] + a['h'], b['y'] + b['h']) - max(a['y'], b['y'])
                if ox > 0.1 and oy > 0.08:
                    area = ox * oy
                    small = min(a['w'] * a['h'], b['w'] * b['h'])
                    if small > 0 and area / small > 0.35:
                        issues.append(f"стр.{a['idx']:>2} наложение текста: "
                                      f"«{a['text'][:28]}» × «{b['text'][:28]}»")
    return issues


if __name__ == '__main__':
    for f in sys.argv[1:]:
        iss = report(f)
        print(f'=== {Path(f).name}: {len(iss)} замечаний')
        for i in iss:
            print('   ', i)
